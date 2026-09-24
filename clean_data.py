import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.utils import get_column_letter

INPUT_FILE = "customer_data.xlsx"
OUTPUT_FILE = "professional_customer_report.xlsx"

print("Reading customer data...")

try:
    df = pd.read_excel(INPUT_FILE)
except FileNotFoundError:
    print("ERROR: customer_data.xlsx not found.")
    raise SystemExit

print("Records loaded:", len(df))

df.columns = df.columns.astype(str).str.strip()

required_columns = [
    "Customer ID",
    "Customer Name",
    "Email",
    "Country",
    "Phone",
    "Order Value"
]

missing_columns = [
    col for col in required_columns
    if col not in df.columns
]

if missing_columns:
    print("ERROR: Required columns are missing:")
    print(missing_columns)
    raise SystemExit

df["Email"] = (
    df["Email"]
    .fillna("")
    .astype(str)
    .str.strip()
    .str.lower()
)

df["Country"] = (
    df["Country"]
    .fillna("Unknown")
    .astype(str)
    .str.strip()
)

df["Phone"] = (
    df["Phone"]
    .fillna("")
    .astype(str)
    .str.strip()
)

df["Duplicate Check"] = (
    df["Email"].ne("")
    & df["Email"].duplicated(keep=False)
)

df["Clean Country"] = df["Country"].replace({
    "USA": "United States",
    "US": "United States",
    "U.S.A": "United States",
    "United States": "United States",
    "UK": "United Kingdom",
    "U.K": "United Kingdom",
    "United Kingdom": "United Kingdom",
    "Canada": "Canada"
})


def check_phone(phone):
    digits = ""

    for character in str(phone):
        if character.isdigit():
            digits += character

    if len(digits) < 10:
        return "Check Phone"

    return "Valid"


df["Phone Quality"] = df["Phone"].apply(check_phone)

duplicate_report = df[
    df["Duplicate Check"] == True
].copy()

phone_report = df[
    df["Phone Quality"] == "Check Phone"
].copy()

total_records = len(df)

duplicate_count = int(
    df["Duplicate Check"].sum()
)

phone_issues = int(
    (df["Phone Quality"] == "Check Phone").sum()
)

valid_phones = int(
    (df["Phone Quality"] == "Valid").sum()
)

unique_countries = int(
    df["Clean Country"].nunique()
)

summary = pd.DataFrame({
    "Metric": [
        "Total Records",
        "Duplicate Email Records",
        "Phone Records Needing Review",
        "Valid Phone Records",
        "Unique Countries"
    ],
    "Value": [
        total_records,
        duplicate_count,
        phone_issues,
        valid_phones,
        unique_countries
    ]
})

country_summary = (
    df["Clean Country"]
    .value_counts()
    .reset_index()
)

country_summary.columns = [
    "Country",
    "Customers"
]

phone_summary = (
    df["Phone Quality"]
    .value_counts()
    .reset_index()
)

phone_summary.columns = [
    "Phone Status",
    "Records"
]

print("Creating Excel report...")

with pd.ExcelWriter(
    OUTPUT_FILE,
    engine="openpyxl"
) as writer:

    summary.to_excel(
        writer,
        sheet_name="Dashboard",
        index=False,
        startrow=3
    )

    df.to_excel(
        writer,
        sheet_name="Clean Customer Data",
        index=False
    )

    duplicate_report.to_excel(
        writer,
        sheet_name="Duplicate Emails",
        index=False
    )

    phone_report.to_excel(
        writer,
        sheet_name="Phone Issues",
        index=False
    )

    country_summary.to_excel(
        writer,
        sheet_name="Country Analysis",
        index=False
    )

    phone_summary.to_excel(
        writer,
        sheet_name="Phone Analysis",
        index=False
    )

print("Applying formatting...")

wb = load_workbook(OUTPUT_FILE)

header_fill = PatternFill(
    fill_type="solid",
    fgColor="1F4E78"
)

header_font = Font(
    bold=True,
    color="FFFFFF"
)

dashboard = wb["Dashboard"]

dashboard["A1"] = "CUSTOMER DATA QUALITY DASHBOARD"
dashboard["A1"].font = Font(
    bold=True,
    size=18
)

dashboard["A1"].alignment = Alignment(
    horizontal="center"
)

dashboard.merge_cells("A1:B1")

dashboard["A2"] = (
    "Python Automated Data Cleaning & Reporting"
)

dashboard["A2"].font = Font(
    italic=True,
    size=11
)

dashboard.merge_cells("A2:B2")

dashboard["A4"] = "Metric"
dashboard["B4"] = "Value"

for cell in dashboard[4]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(
        horizontal="center"
    )

for row in range(5, 10):
    dashboard[f"A{row}"].font = Font(
        bold=True
    )

    dashboard[f"B{row}"].alignment = Alignment(
        horizontal="center"
    )

dashboard.column_dimensions["A"].width = 35
dashboard.column_dimensions["B"].width = 20

for sheet_name in wb.sheetnames:

    worksheet = wb[sheet_name]

    worksheet.freeze_panes = "A2"

    if sheet_name == "Dashboard":
        continue

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center"
        )

    for column_cells in worksheet.columns:

        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:

            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(max_length + 2, 12),
            40
        )

country_sheet = wb["Country Analysis"]

bar_chart = BarChart()

bar_chart.title = "Customers by Country"
bar_chart.y_axis.title = "Number of Customers"
bar_chart.x_axis.title = "Country"

bar_chart.height = 7
bar_chart.width = 12

country_data = Reference(
    country_sheet,
    min_col=2,
    min_row=1,
    max_row=country_sheet.max_row
)

country_categories = Reference(
    country_sheet,
    min_col=1,
    min_row=2,
    max_row=country_sheet.max_row
)

bar_chart.add_data(
    country_data,
    titles_from_data=True
)

bar_chart.set_categories(
    country_categories
)

dashboard.add_chart(
    bar_chart,
    "D4"
)

phone_analysis_sheet = wb["Phone Analysis"]

pie_chart = PieChart()

pie_chart.title = "Phone Data Quality"

pie_chart.height = 7
pie_chart.width = 10

phone_data = Reference(
    phone_analysis_sheet,
    min_col=2,
    min_row=1,
    max_row=phone_analysis_sheet.max_row
)

phone_categories = Reference(
    phone_analysis_sheet,
    min_col=1,
    min_row=2,
    max_row=phone_analysis_sheet.max_row
)

pie_chart.add_data(
    phone_data,
    titles_from_data=True
)

pie_chart.set_categories(
    phone_categories
)

dashboard.add_chart(
    pie_chart,
    "D20"
)

try:
    wb.save(OUTPUT_FILE)
except PermissionError:
    print("Please close the Excel output file and run again.")
    raise SystemExit

print()
print("========================================")
print(" PROFESSIONAL DASHBOARD CREATED")
print("========================================")
print("Total Records:", total_records)
print("Duplicate Emails:", duplicate_count)
print("Phone Issues:", phone_issues)
print("Valid Phones:", valid_phones)
print("Unique Countries:", unique_countries)
print("File:", OUTPUT_FILE)
print("DONE!")